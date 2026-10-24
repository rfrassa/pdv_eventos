"""
generar_reporte_evento.py
Extrae TODOS los datos de un evento y genera HTML + Excel.
Uso: python generar_reporte_evento.py [evento_id]
     Si no se pasa evento_id, usa el evento activo.
"""
import os
import sys

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'ibat_pdv_eventos.settings')
import django
django.setup()

from collections import defaultdict
from decimal import Decimal
from datetime import datetime

from django.db.models import Sum, Count, Q
from pdv.models import Evento, PuntoVenta, Categoria, Producto, Pedido, LineaPedido, Pago

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl no instalado. Ejecutá: pip install openpyxl")
    sys.exit(1)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

METODOS = {'EF': 'Efectivo', 'TC': 'T. Crédito', 'TD': 'T. Débito',
           'TR': 'Transferencia/QR', 'CU': 'Cta. Corriente', 'OT': 'Otro'}


def fmt_peso(v):
    return f"${v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def get_evento(evento_id=None):
    if evento_id:
        return Evento.objects.get(pk=evento_id)
    ev = Evento.objects.filter(activo=True).first()
    if not ev:
        ev = Evento.objects.order_by('-fecha').first()
    return ev


# ---------------------------------------------------------------------------
# Extracción de datos
# ---------------------------------------------------------------------------

def extraer_datos(evento):
    pedidos = Pedido.objects.filter(
        evento=evento, cerrado=True
    ).select_related('punto_venta').prefetch_related('lineas__producto__categoria', 'pagos')

    total_general = Decimal('0')
    total_impuestos = Decimal('0')
    total_descuentos = Decimal('0')
    pedidos_con_descuento = 0

    por_metodo = defaultdict(Decimal)
    por_producto = defaultdict(lambda: {'unidades': 0, 'revenue': Decimal('0'),
                                        'categoria': '', 'precio': Decimal('0')})
    por_categoria = defaultdict(lambda: {'unidades': 0, 'revenue': Decimal('0')})
    por_pdv = defaultdict(lambda: {'pedidos': 0, 'total': Decimal('0')})
    por_hora = defaultdict(lambda: {'pedidos': 0, 'total': Decimal('0')})
    notas_especiales = []
    detalle_pedidos = []
    detalle_lineas = []

    for p in pedidos:
        total_general += p.total_final
        total_impuestos += p.total_impuestos
        descuento_monto = (p.subtotal * p.descuento_porcentaje / 100).quantize(Decimal('0.01'))
        total_descuentos += descuento_monto
        if p.descuento_porcentaje > 0:
            pedidos_con_descuento += 1

        por_pdv[p.punto_venta.nombre]['pedidos'] += 1
        por_pdv[p.punto_venta.nombre]['total'] += p.total_final

        hora = p.creado.hour
        por_hora[hora]['pedidos'] += 1
        por_hora[hora]['total'] += p.total_final

        metodos_pedido = []
        for pago in p.pagos.all():
            por_metodo[pago.metodo] += pago.monto
            metodos_pedido.append(f"{METODOS.get(pago.metodo, pago.metodo)} {fmt_peso(pago.monto)}")

        detalle_pedidos.append({
            'id': p.id,
            'hora': p.creado.strftime('%H:%M:%S'),
            'pdv': p.punto_venta.nombre,
            'subtotal': p.subtotal,
            'descuento_pct': p.descuento_porcentaje,
            'descuento_monto': descuento_monto,
            'impuestos': p.total_impuestos,
            'total': p.total_final,
            'metodos': ' | '.join(metodos_pedido),
            'veces_impreso': p.veces_impreso,
        })

        for linea in p.lineas.all():
            prod = linea.producto
            cat_nombre = prod.categoria.nombre if prod.categoria else '(sin categoría)'
            subtotal_linea = linea.precio_unitario * linea.cantidad

            por_producto[prod.nombre]['unidades'] += linea.cantidad
            por_producto[prod.nombre]['revenue'] += subtotal_linea
            por_producto[prod.nombre]['categoria'] = cat_nombre
            por_producto[prod.nombre]['precio'] = linea.precio_unitario

            por_categoria[cat_nombre]['unidades'] += linea.cantidad
            por_categoria[cat_nombre]['revenue'] += subtotal_linea

            if linea.nota:
                notas_especiales.append({
                    'pedido_id': p.id,
                    'hora': p.creado.strftime('%H:%M'),
                    'producto': prod.nombre,
                    'cantidad': linea.cantidad,
                    'nota': linea.nota,
                })

            detalle_lineas.append({
                'pedido_id': p.id,
                'hora': p.creado.strftime('%H:%M:%S'),
                'pdv': p.punto_venta.nombre,
                'categoria': cat_nombre,
                'producto': prod.nombre,
                'cantidad': linea.cantidad,
                'precio_unitario': linea.precio_unitario,
                'subtotal': subtotal_linea,
                'nota': linea.nota or '',
            })

    # Productos disponibles sin ventas
    todos_productos = Producto.objects.filter(evento=evento, disponible=True)
    sin_ventas = [p.nombre for p in todos_productos if p.nombre not in por_producto]

    # Hora pico
    hora_pico = max(por_hora.items(), key=lambda x: x[1]['total'], default=(None, {}))

    return {
        'evento': evento,
        'total_general': total_general,
        'total_impuestos': total_impuestos,
        'total_descuentos': total_descuentos,
        'total_pedidos': len(detalle_pedidos),
        'pedidos_con_descuento': pedidos_con_descuento,
        'ticket_promedio': (total_general / len(detalle_pedidos)).quantize(Decimal('0.01')) if detalle_pedidos else Decimal('0'),
        'por_metodo': dict(sorted(por_metodo.items(), key=lambda x: -x[1])),
        'por_producto': dict(sorted(por_producto.items(), key=lambda x: -x[1]['revenue'])),
        'por_categoria': dict(sorted(por_categoria.items(), key=lambda x: -x[1]['revenue'])),
        'por_pdv': dict(por_pdv),
        'por_hora': dict(sorted(por_hora.items())),
        'hora_pico': hora_pico[0],
        'notas_especiales': notas_especiales,
        'sin_ventas': sin_ventas,
        'detalle_pedidos': sorted(detalle_pedidos, key=lambda x: x['id']),
        'detalle_lineas': sorted(detalle_lineas, key=lambda x: x['pedido_id']),
    }


# ---------------------------------------------------------------------------
# Generador HTML
# ---------------------------------------------------------------------------

def generar_html(d, ruta_salida):
    ev = d['evento']
    fecha_str = ev.fecha.strftime('%d/%m/%Y')
    ahora = datetime.now().strftime('%d/%m/%Y %H:%M')

    def tabla_kv(filas):
        html = '<table class="kv">'
        for k, v in filas:
            html += f'<tr><td class="label">{k}</td><td class="value">{v}</td></tr>'
        html += '</table>'
        return html

    # Sección: resumen ejecutivo
    sec_resumen = f"""
    <div class="section">
      <h2>Resumen ejecutivo</h2>
      {tabla_kv([
        ('Total vendido', f'<strong>{fmt_peso(d["total_general"])}</strong>'),
        ('Pedidos cerrados', str(d['total_pedidos'])),
        ('Ticket promedio', fmt_peso(d['ticket_promedio'])),
        ('Total impuestos', fmt_peso(d['total_impuestos'])),
        ('Total descuentos', fmt_peso(d['total_descuentos'])),
        ('Pedidos con descuento', str(d['pedidos_con_descuento'])),
        ('Hora pico', f'{d["hora_pico"]:02d}:00 — {d["hora_pico"]+1:02d}:00' if d['hora_pico'] is not None else '—'),
        ('Productos sin ventas', str(len(d['sin_ventas']))),
      ])}
    </div>"""

    # Sección: por método de pago
    rows_met = ''.join(
        f'<tr><td>{METODOS.get(m, m)}</td><td class="num">{fmt_peso(t)}</td>'
        f'<td class="num">{t/d["total_general"]*100:.1f}%</td></tr>'
        for m, t in d['por_metodo'].items()
    )
    sec_metodo = f"""
    <div class="section">
      <h2>Por método de pago</h2>
      <table class="data">
        <thead><tr><th>Método</th><th>Total</th><th>%</th></tr></thead>
        <tbody>{rows_met}</tbody>
      </table>
    </div>"""

    # Sección: por categoría
    rows_cat = ''.join(
        f'<tr><td>{c}</td><td class="num">{v["unidades"]}</td>'
        f'<td class="num">{fmt_peso(v["revenue"])}</td></tr>'
        for c, v in d['por_categoria'].items()
    )
    sec_cat = f"""
    <div class="section">
      <h2>Por categoría</h2>
      <table class="data">
        <thead><tr><th>Categoría</th><th>Unidades</th><th>Revenue</th></tr></thead>
        <tbody>{rows_cat}</tbody>
      </table>
    </div>"""

    # Sección: por producto
    rows_prod = ''.join(
        f'<tr><td>{n}</td><td>{v["categoria"]}</td>'
        f'<td class="num">{v["unidades"]}</td>'
        f'<td class="num">{fmt_peso(v["revenue"])}</td>'
        f'<td class="num">{fmt_peso(v["precio"])}</td></tr>'
        for n, v in d['por_producto'].items()
    )
    sec_prod = f"""
    <div class="section">
      <h2>Por producto (ordenado por revenue)</h2>
      <table class="data">
        <thead><tr><th>Producto</th><th>Categoría</th><th>Unidades</th><th>Revenue</th><th>Precio unit.</th></tr></thead>
        <tbody>{rows_prod}</tbody>
      </table>
    </div>"""

    # Sección: por PDV
    rows_pdv = ''.join(
        f'<tr><td>{n}</td><td class="num">{v["pedidos"]}</td>'
        f'<td class="num">{fmt_peso(v["total"])}</td></tr>'
        for n, v in d['por_pdv'].items()
    )
    sec_pdv = f"""
    <div class="section">
      <h2>Por punto de venta</h2>
      <table class="data">
        <thead><tr><th>Punto de venta</th><th>Pedidos</th><th>Total</th></tr></thead>
        <tbody>{rows_pdv}</tbody>
      </table>
    </div>"""

    # Sección: ventas por hora
    rows_hora = ''.join(
        f'<tr><td>{h:02d}:00 — {h+1:02d}:00</td>'
        f'<td class="num">{v["pedidos"]}</td>'
        f'<td class="num">{fmt_peso(v["total"])}</td></tr>'
        for h, v in d['por_hora'].items()
    )
    sec_hora = f"""
    <div class="section">
      <h2>Ventas por hora</h2>
      <table class="data">
        <thead><tr><th>Franja</th><th>Pedidos</th><th>Total</th></tr></thead>
        <tbody>{rows_hora}</tbody>
      </table>
    </div>"""

    # Sección: notas especiales
    if d['notas_especiales']:
        rows_notas = ''.join(
            f'<tr><td>#{n["pedido_id"]}</td><td>{n["hora"]}</td>'
            f'<td>{n["producto"]}</td><td class="num">{n["cantidad"]}</td>'
            f'<td>{n["nota"]}</td></tr>'
            for n in d['notas_especiales']
        )
        sec_notas = f"""
        <div class="section">
          <h2>Notas especiales en pedidos ({len(d["notas_especiales"])})</h2>
          <table class="data">
            <thead><tr><th>Pedido</th><th>Hora</th><th>Producto</th><th>Cant.</th><th>Nota</th></tr></thead>
            <tbody>{rows_notas}</tbody>
          </table>
        </div>"""
    else:
        sec_notas = ''

    # Sección: productos sin ventas
    if d['sin_ventas']:
        items_sv = ''.join(f'<li>{p}</li>' for p in sorted(d['sin_ventas']))
        sec_sin_ventas = f"""
        <div class="section">
          <h2>Productos disponibles sin ventas ({len(d["sin_ventas"])})</h2>
          <ul class="sin-ventas">{items_sv}</ul>
        </div>"""
    else:
        sec_sin_ventas = '<div class="section"><h2>Productos sin ventas</h2><p>Todos los productos tuvieron al menos una venta.</p></div>'

    # Sección: detalle de pedidos
    rows_det = ''.join(
        f'<tr><td class="num">#{p["id"]}</td><td>{p["hora"]}</td>'
        f'<td>{p["pdv"]}</td>'
        f'<td class="num">{fmt_peso(p["subtotal"])}</td>'
        f'<td class="num">{p["descuento_pct"]:.0f}%</td>'
        f'<td class="num">{fmt_peso(p["impuestos"])}</td>'
        f'<td class="num">{fmt_peso(p["total"])}</td>'
        f'<td>{p["metodos"]}</td></tr>'
        for p in d['detalle_pedidos']
    )
    sec_detalle = f"""
    <div class="section pagebreak">
      <h2>Detalle de pedidos ({d["total_pedidos"]} órdenes)</h2>
      <table class="data small">
        <thead><tr><th>ID</th><th>Hora</th><th>PDV</th><th>Subtotal</th><th>Dto%</th><th>IVA</th><th>Total</th><th>Métodos</th></tr></thead>
        <tbody>{rows_det}</tbody>
      </table>
    </div>"""

    html = f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<title>Reporte {ev.nombre} — {fecha_str}</title>
<style>
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: Arial, sans-serif; font-size: 13px; color: #1a1a1a; background: #f5f5f5; }}
  .cover {{ background: #1a237e; color: white; padding: 48px 40px 40px; margin-bottom: 24px; }}
  .cover h1 {{ font-size: 28px; margin-bottom: 8px; }}
  .cover .sub {{ font-size: 14px; opacity: .8; }}
  .container {{ max-width: 1100px; margin: 0 auto; padding: 0 16px 48px; }}
  .section {{ background: white; border-radius: 6px; padding: 20px 24px; margin-bottom: 18px;
              box-shadow: 0 1px 3px rgba(0,0,0,.12); }}
  h2 {{ font-size: 15px; color: #1a237e; border-bottom: 2px solid #e8eaf6;
        padding-bottom: 8px; margin-bottom: 14px; }}
  table.kv {{ width: 100%; border-collapse: collapse; }}
  table.kv td {{ padding: 6px 8px; border-bottom: 1px solid #f0f0f0; }}
  table.kv .label {{ color: #555; width: 220px; }}
  table.kv .value {{ font-weight: 500; }}
  table.data {{ width: 100%; border-collapse: collapse; font-size: 12px; }}
  table.data.small {{ font-size: 11px; }}
  table.data thead tr {{ background: #1a237e; color: white; }}
  table.data th {{ padding: 7px 10px; text-align: left; }}
  table.data td {{ padding: 5px 10px; border-bottom: 1px solid #f0f0f0; }}
  table.data tbody tr:hover {{ background: #f5f5ff; }}
  .num {{ text-align: right; font-variant-numeric: tabular-nums; }}
  .sin-ventas {{ columns: 3; padding-left: 18px; }}
  .sin-ventas li {{ margin-bottom: 4px; }}
  .footer {{ text-align: center; color: #999; font-size: 11px; margin-top: 24px; }}
  @media print {{
    body {{ background: white; }}
    .section {{ box-shadow: none; border: 1px solid #ddd; }}
    .pagebreak {{ page-break-before: always; }}
  }}
</style>
</head>
<body>
<div class="cover">
  <h1>Reporte de Evento — {ev.nombre}</h1>
  <div class="sub">{fecha_str} &nbsp;·&nbsp; Generado el {ahora}</div>
</div>
<div class="container">
  {sec_resumen}
  {sec_metodo}
  {sec_cat}
  {sec_prod}
  {sec_pdv}
  {sec_hora}
  {sec_notas}
  {sec_sin_ventas}
  {sec_detalle}
</div>
<div class="footer">IBAT San José — Sistema PDV Eventos · {ahora}</div>
</body>
</html>"""

    with open(ruta_salida, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f"  HTML generado: {ruta_salida}")


# ---------------------------------------------------------------------------
# Generador Excel
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill('solid', fgColor='1A237E')
HEADER_FONT = Font(color='FFFFFF', bold=True, size=10)
ALT_FILL = PatternFill('solid', fgColor='E8EAF6')
BOLD = Font(bold=True)
TOTAL_FILL = PatternFill('solid', fgColor='C5CAE9')


def _header(ws, cols):
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center')


def _autofit(ws, min_w=10, max_w=50):
    for col in ws.columns:
        length = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(length + 2, min_w), max_w)


def _freeze(ws, cell='A2'):
    ws.freeze_panes = cell


def generar_excel(d, ruta_salida):
    wb = openpyxl.Workbook()
    ev = d['evento']
    fecha_str = ev.fecha.strftime('%d/%m/%Y')

    # --- Hoja 1: Resumen ---
    ws = wb.active
    ws.title = 'Resumen'
    ws.append(['REPORTE DE EVENTO'])
    ws['A1'].font = Font(bold=True, size=14, color='1A237E')
    ws.append([ev.nombre, '', fecha_str])
    ws.append([])
    _header(ws, ['Indicador', 'Valor'])
    kpis = [
        ('Total vendido', float(d['total_general'])),
        ('Pedidos cerrados', d['total_pedidos']),
        ('Ticket promedio', float(d['ticket_promedio'])),
        ('Total impuestos', float(d['total_impuestos'])),
        ('Total descuentos aplicados', float(d['total_descuentos'])),
        ('Pedidos con descuento', d['pedidos_con_descuento']),
        ('Hora pico', f'{d["hora_pico"]:02d}:00' if d['hora_pico'] is not None else '—'),
        ('Productos sin ventas', len(d['sin_ventas'])),
    ]
    for i, (k, v) in enumerate(kpis):
        ws.append([k, v])
        if i % 2 == 0:
            ws[f'A{ws.max_row}'].fill = ALT_FILL
            ws[f'B{ws.max_row}'].fill = ALT_FILL
    _autofit(ws)
    _freeze(ws)

    # --- Hoja 2: Métodos de pago ---
    ws2 = wb.create_sheet('Métodos de pago')
    _header(ws2, ['Método', 'Total ($)', '% del total'])
    for m, t in d['por_metodo'].items():
        pct = float(t / d['total_general'] * 100) if d['total_general'] else 0
        ws2.append([METODOS.get(m, m), float(t), round(pct, 2)])
    ws2.append([])
    ws2.append(['TOTAL', float(d['total_general']), 100.0])
    ws2[f'A{ws2.max_row}'].font = BOLD
    ws2[f'B{ws2.max_row}'].font = BOLD
    _autofit(ws2)

    # --- Hoja 3: Por producto ---
    ws3 = wb.create_sheet('Por producto')
    _header(ws3, ['Producto', 'Categoría', 'Unidades', 'Revenue ($)', 'Precio unit. ($)', '% revenue'])
    for n, v in d['por_producto'].items():
        pct = float(v['revenue'] / d['total_general'] * 100) if d['total_general'] else 0
        ws3.append([n, v['categoria'], v['unidades'], float(v['revenue']),
                    float(v['precio']), round(pct, 2)])
    _autofit(ws3)
    _freeze(ws3)

    # --- Hoja 4: Por categoría ---
    ws4 = wb.create_sheet('Por categoría')
    _header(ws4, ['Categoría', 'Unidades', 'Revenue ($)', '% revenue'])
    for c, v in d['por_categoria'].items():
        pct = float(v['revenue'] / d['total_general'] * 100) if d['total_general'] else 0
        ws4.append([c, v['unidades'], float(v['revenue']), round(pct, 2)])
    _autofit(ws4)

    # --- Hoja 5: Por punto de venta ---
    ws5 = wb.create_sheet('Por PDV')
    _header(ws5, ['Punto de venta', 'Pedidos', 'Total ($)'])
    for n, v in d['por_pdv'].items():
        ws5.append([n, v['pedidos'], float(v['total'])])
    _autofit(ws5)

    # --- Hoja 6: Ventas por hora ---
    ws6 = wb.create_sheet('Por hora')
    _header(ws6, ['Franja horaria', 'Pedidos', 'Total ($)'])
    for h, v in d['por_hora'].items():
        ws6.append([f'{h:02d}:00 — {h+1:02d}:00', v['pedidos'], float(v['total'])])
    _autofit(ws6)

    # --- Hoja 7: Detalle pedidos ---
    ws7 = wb.create_sheet('Detalle pedidos')
    _header(ws7, ['ID', 'Hora', 'Punto de venta', 'Subtotal ($)',
                  'Dto (%)', 'Dto ($)', 'Impuestos ($)', 'Total ($)',
                  'Métodos de pago', 'Veces impreso'])
    for p in d['detalle_pedidos']:
        ws7.append([p['id'], p['hora'], p['pdv'],
                    float(p['subtotal']), float(p['descuento_pct']),
                    float(p['descuento_monto']), float(p['impuestos']),
                    float(p['total']), p['metodos'], p['veces_impreso']])
    _autofit(ws7)
    _freeze(ws7)

    # --- Hoja 8: Líneas completas (auditoría) ---
    ws8 = wb.create_sheet('Líneas completas')
    _header(ws8, ['Pedido ID', 'Hora', 'PDV', 'Categoría', 'Producto',
                  'Cantidad', 'Precio unit. ($)', 'Subtotal ($)', 'Nota'])
    for l in d['detalle_lineas']:
        ws8.append([l['pedido_id'], l['hora'], l['pdv'], l['categoria'],
                    l['producto'], l['cantidad'],
                    float(l['precio_unitario']), float(l['subtotal']), l['nota']])
    _autofit(ws8)
    _freeze(ws8)

    # --- Hoja 9: Productos sin ventas ---
    ws9 = wb.create_sheet('Sin ventas')
    _header(ws9, ['Producto (disponible, 0 ventas)'])
    for p in sorted(d['sin_ventas']):
        ws9.append([p])
    _autofit(ws9)

    # --- Hoja 10: Notas especiales ---
    ws10 = wb.create_sheet('Notas especiales')
    _header(ws10, ['Pedido ID', 'Hora', 'Producto', 'Cantidad', 'Nota'])
    for n in d['notas_especiales']:
        ws10.append([n['pedido_id'], n['hora'], n['producto'], n['cantidad'], n['nota']])
    _autofit(ws10)

    wb.save(ruta_salida)
    print(f"  Excel generado:  {ruta_salida}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

if __name__ == '__main__':
    evento_id = int(sys.argv[1]) if len(sys.argv) > 1 else None

    print("Extrayendo datos del evento...")
    evento = get_evento(evento_id)
    if not evento:
        print("ERROR: No se encontró ningún evento.")
        sys.exit(1)

    print(f"  Evento: {evento.nombre} ({evento.fecha})")
    datos = extraer_datos(evento)
    print(f"  Pedidos procesados: {datos['total_pedidos']}")
    print(f"  Total vendido:      {fmt_peso(datos['total_general'])}")
    print(f"  Líneas de detalle:  {len(datos['detalle_lineas'])}")
    print()

    slug = evento.nombre.lower().replace(' ', '_').replace('/', '-')
    fecha_slug = evento.fecha.strftime('%Y-%m-%d')
    base = f"reporte_{slug}_{fecha_slug}"

    print("Generando archivos...")
    generar_html(datos, f"{base}.html")
    generar_excel(datos, f"{base}.xlsx")
    print()
    print("Listo.")
