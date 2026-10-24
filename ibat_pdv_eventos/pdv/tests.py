import re
from datetime import date, datetime, timezone as dtz
from decimal import Decimal
from io import BytesIO, StringIO

from unittest import mock

from django.contrib.auth.models import Group, User
from django.core.management import call_command
from django.db import IntegrityError, transaction
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from .models import Categoria, Evento, LineaPedido, Pago, Pedido, Producto, PuntoVenta
from .views import resumen_ventas


def utc(*args):
    return datetime(*args, tzinfo=dtz.utc)


class DatosVentas(TestCase):
    """Base sin tests: dos eventos, dos cajas y dos productos."""

    @classmethod
    def setUpTestData(cls):
        cls.junio = Evento.objects.create(nombre='Peña 2026', fecha=date(2026, 6, 12), activo=False)
        cls.expo = Evento.objects.create(nombre='Expo 2026', fecha=date(2026, 10, 22), activo=True)
        # La caja se usó en junio y después se pasó a la Expo (como en producción).
        cls.caja = PuntoVenta.objects.create(nombre='Caja 1', evento=cls.expo)
        cls.caja2 = PuntoVenta.objects.create(nombre='Caja 2', evento=cls.expo)
        cat = Categoria.objects.create(nombre='Comidas', evento=cls.expo)
        cls.choripan = Producto.objects.create(nombre='Choripán', precio=Decimal('1000'), categoria=cat, evento=cls.expo)
        cls.agua = Producto.objects.create(nombre='Agua', precio=Decimal('500'), categoria=cat, evento=cls.expo)

    def pedido(self, evento, creado, cantidad=1, producto=None, metodo='EF', cerrado=True, caja=None):
        producto = producto or self.choripan
        total = producto.precio * cantidad
        p = Pedido.objects.create(
            punto_venta=caja or self.caja, evento=evento, cerrado=cerrado,
            subtotal=total, total_final=total,
        )
        Pedido.objects.filter(id=p.id).update(creado=creado)  # auto_now_add
        LineaPedido.objects.create(pedido=p, producto=producto, cantidad=cantidad, precio_unitario=producto.precio)
        if cerrado:
            Pago.objects.create(pedido=p, metodo=metodo, monto=total)
        return p


class ResumenVentasTests(DatosVentas):
    """El resumen es el número que se le informa al colegio: cada evento por separado,
    solo pedidos cobrados y fechas en hora de Argentina (UTC-3)."""

    def test_cada_evento_por_separado_y_sin_pedidos_sin_cobrar(self):
        self.pedido(self.junio, utc(2026, 6, 12, 15, 0), cantidad=2)            # 2.000
        self.pedido(self.junio, utc(2026, 6, 12, 16, 0), producto=self.agua)    # 500
        self.pedido(self.expo, utc(2026, 10, 23, 0, 30), cantidad=3)            # 3.000
        self.pedido(self.expo, utc(2026, 10, 23, 0, 40), cerrado=False)         # abierto: no cuenta

        r_junio = resumen_ventas(self.junio)
        r_expo = resumen_ventas(self.expo)

        self.assertEqual((r_junio['total_pedidos'], r_junio['total_general']), (2, 2500))
        self.assertEqual((r_expo['total_pedidos'], r_expo['total_general']), (1, 3000))
        self.assertEqual(r_junio['ticket_promedio'], 1250)
        self.assertEqual(r_junio['unidades_total'], 3)
        self.assertEqual([(p['nombre'], p['unidades'], p['importe']) for p in r_junio['por_producto']],
                         [('Choripán', 2, 2000), ('Agua', 1, 500)])
        self.assertEqual(r_expo['total_por_metodo'][0]['total'], 3000)

    def test_todos_los_eventos(self):
        self.pedido(self.junio, utc(2026, 6, 12, 15, 0))
        self.pedido(self.expo, utc(2026, 10, 23, 0, 30))
        r = resumen_ventas(None)
        self.assertEqual((r['evento'], r['total_pedidos'], r['total_general']), ('Todos los eventos', 2, 2000))

    def test_se_cuenta_por_el_evento_del_pedido_no_por_el_de_la_caja(self):
        # self.caja hoy pertenece a la Expo, pero este pedido es de junio.
        self.pedido(self.junio, utc(2026, 6, 12, 15, 0))
        self.assertEqual(resumen_ventas(self.junio)['total_pedidos'], 1)
        self.assertEqual(resumen_ventas(self.expo)['total_pedidos'], 0)

    def test_rango_de_fechas_en_hora_de_argentina(self):
        antes_de_medianoche = self.pedido(self.junio, utc(2026, 6, 13, 2, 30))   # 12/06 23:30 ART
        despues_de_medianoche = self.pedido(self.junio, utc(2026, 6, 13, 3, 30))  # 13/06 00:30 ART

        dia12 = resumen_ventas(self.junio, desde=date(2026, 6, 12), hasta=date(2026, 6, 12))
        dia13 = resumen_ventas(self.junio, desde=date(2026, 6, 13), hasta=date(2026, 6, 13))
        self.assertEqual(dia12['total_pedidos'], 1)
        self.assertEqual(dia13['total_pedidos'], 1)
        self.assertEqual(resumen_ventas(self.junio, desde=date(2026, 6, 12), hasta=date(2026, 6, 13))['total_pedidos'], 2)
        self.assertEqual(resumen_ventas(self.junio, desde=date(2026, 6, 13))['total_pedidos'], 1)
        self.assertEqual(resumen_ventas(self.junio, hasta=date(2026, 6, 12))['total_pedidos'], 1)
        self.assertIsNotNone(antes_de_medianoche.id)
        self.assertIsNotNone(despues_de_medianoche.id)

    def test_franjas_de_15_minutos_en_hora_local(self):
        self.pedido(self.junio, utc(2026, 6, 12, 13, 7))    # 10:07 ART
        self.pedido(self.junio, utc(2026, 6, 12, 13, 14))   # 10:14 ART
        self.pedido(self.junio, utc(2026, 6, 12, 13, 15))   # 10:15 ART
        franjas = resumen_ventas(self.junio)['por_franja']
        self.assertEqual([(f['franja'], f['pedidos'], f['total']) for f in franjas],
                         [('12/06 10:00', 2, 2000), ('12/06 10:15', 1, 1000)])

    def test_por_caja_y_metodo(self):
        self.pedido(self.expo, utc(2026, 10, 23, 0, 30), metodo='EF')
        self.pedido(self.expo, utc(2026, 10, 23, 0, 31), metodo='TR', caja=self.caja2)
        cajas = {c['nombre']: c for c in resumen_ventas(self.expo)['por_caja']}
        self.assertEqual(cajas['Caja 1']['por_metodo'], {'EF': 1000})
        self.assertEqual(cajas['Caja 2']['por_metodo'], {'TR': 1000})
        self.assertEqual(sum(c['total'] for c in cajas.values()), 2000)


class TableroEndpointsTests(DatosVentas):
    """Las tres salidas (JSON, XLSX, HTML) salen del mismo cálculo."""

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(User.objects.create_superuser('admin', 'a@a.com', 'x'))
        self.pedido(self.junio, utc(2026, 6, 12, 15, 0), cantidad=1500)   # 1.500.000
        self.pedido(self.expo, utc(2026, 10, 23, 0, 30), cantidad=3)
        self.pedido(self.expo, utc(2026, 10, 23, 0, 40), cerrado=False)

    def test_por_defecto_es_el_evento_activo(self):
        data = self.client.get('/api/resumen-ventas/').json()
        self.assertEqual((data['evento_id'], data['total_pedidos'], data['total_general']), (self.expo.id, 1, 3000))

    def test_evento_explicito_y_todos(self):
        data = self.client.get('/api/resumen-ventas/', {'evento': self.junio.id}).json()
        self.assertEqual(data['total_general'], 1500000)
        data = self.client.get('/api/resumen-ventas/', {'evento': 'todos'}).json()
        self.assertEqual((data['total_pedidos'], data['total_general']), (2, 1503000))

    def test_parametros_invalidos_dan_400(self):
        for params in ({'desde': '23-10-2026'}, {'evento': 'abc'}, {'punto_venta_id': 'x'},
                       {'desde': '2026-10-23', 'hasta': '2026-10-22'}):
            self.assertEqual(self.client.get('/api/resumen-ventas/', params).status_code, 400, params)
        self.assertEqual(self.client.get('/api/resumen-ventas/', {'evento': 9999}).status_code, 404)

    def test_xlsx_tiene_las_hojas_y_los_mismos_totales(self):
        resp = self.client.get('/api/resumen-ventas/xlsx/', {'evento': self.junio.id})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('attachment; filename="ventas_pena-2026.xlsx"', resp['Content-Disposition'])
        wb = load_workbook(BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, ['Resumen', 'Por producto', 'Por caja y método de pago', 'Pedidos'])
        resumen = {fila[0]: fila[1] for fila in wb['Resumen'].iter_rows(min_row=3, max_row=9, values_only=True)}
        self.assertEqual(resumen['Total vendido'], 1500000)
        self.assertEqual(resumen['Pedidos cobrados'], 1)
        self.assertEqual(wb['Pedidos'].max_row, 2)   # encabezado + 1 pedido

    def test_imprimible_muestra_el_total_formateado(self):
        resp = self.client.get('/api/resumen-ventas/imprimir/', {'evento': self.junio.id})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, '$1.500.000,00')
        self.assertContains(resp, '@media print')


class CierreDeEventoTests(DatosVentas):
    """Cerrar un evento: acción del admin (solo staff), sin tocar pedidos, y con las
    reglas de borrado/edición de pedidos cobrados o de eventos cerrados."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        self.cobrado = self.pedido(self.expo, utc(2026, 10, 23, 0, 30), cantidad=3)
        self.abierto = self.pedido(self.expo, utc(2026, 10, 23, 0, 40), cerrado=False)
        admin = get_user_model().objects.create_superuser('admin', 'a@a.com', 'x')
        self.api = APIClient()
        self.api.force_authenticate(admin)
        self.admin_client = self.client_class()
        self.admin_client.force_login(admin)

    def accion_cerrar(self, cliente, ids, confirmar=False):
        datos = {'action': 'cerrar_evento', '_selected_action': ids}
        if confirmar:
            datos.update({'confirmar': 'si', 'post': 'yes'})
        return cliente.post('/admin/pdv/evento/', datos)

    def cerrar_expo(self):
        self.accion_cerrar(self.admin_client, [self.expo.id], confirmar=True)
        self.expo.refresh_from_db()

    # --- la acción del admin ---
    def test_la_confirmacion_muestra_el_total_y_todavia_no_cierra(self):
        resp = self.accion_cerrar(self.admin_client, [self.expo.id])
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, '$3.000,00')
        self.assertContains(resp, '1 pedido sin cobrar')
        self.expo.refresh_from_db()
        self.assertIsNone(self.expo.cerrado_en)
        self.assertTrue(self.expo.activo)

    def test_confirmar_cierra_sin_tocar_los_pedidos(self):
        antes = list(Pedido.objects.order_by('id').values_list('id', 'cerrado', 'total_final'))
        resp = self.accion_cerrar(self.admin_client, [self.expo.id], confirmar=True)
        self.assertEqual(resp.status_code, 302)
        self.expo.refresh_from_db()
        self.assertFalse(self.expo.activo)
        self.assertIsNotNone(self.expo.cerrado_en)
        self.assertEqual(list(Pedido.objects.order_by('id').values_list('id', 'cerrado', 'total_final')), antes)
        self.assertEqual(resumen_ventas(self.expo)['total_general'], 3000)  # el resumen sigue igual

    def test_sin_login_de_staff_no_cierra(self):
        resp = self.accion_cerrar(self.client_class(), [self.expo.id], confirmar=True)
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/admin/login/', resp['Location'])
        self.expo.refresh_from_db()
        self.assertIsNone(self.expo.cerrado_en)

    def test_no_se_cierran_dos_eventos_a_la_vez(self):
        self.accion_cerrar(self.admin_client, [self.expo.id, self.junio.id], confirmar=True)
        self.expo.refresh_from_db()
        self.junio.refresh_from_db()
        self.assertIsNone(self.expo.cerrado_en)
        self.assertIsNone(self.junio.cerrado_en)

    def test_cerrar_dos_veces_no_cambia_la_hora(self):
        self.cerrar_expo()
        primera = self.expo.cerrado_en
        self.accion_cerrar(self.admin_client, [self.expo.id], confirmar=True)
        self.expo.refresh_from_db()
        self.assertEqual(self.expo.cerrado_en, primera)

    def test_un_evento_cerrado_no_puede_quedar_activo(self):
        from django.core.exceptions import ValidationError
        self.cerrar_expo()
        self.expo.activo = True
        with self.assertRaises(ValidationError):
            self.expo.full_clean()
        self.expo.cerrado_en = None  # reabrir es vaciar "Cerrado en"
        self.expo.full_clean()

    # --- pedidos nuevos ---
    def nuevo_pedido(self):
        return self.api.post('/api/pedidos/', {
            'punto_venta': self.caja.id, 'lineas': [{'producto': self.choripan.id, 'cantidad': 1}],
        }, format='json')

    def test_se_rechazan_pedidos_nuevos_de_un_evento_cerrado(self):
        self.assertEqual(self.nuevo_pedido().status_code, 201)
        self.cerrar_expo()
        resp = self.nuevo_pedido()
        self.assertEqual(resp.status_code, 400)
        self.assertIn('punto_venta', resp.json())

    def test_cerrado_en_rechaza_aunque_quede_activo(self):
        Evento.objects.filter(id=self.expo.id).update(cerrado_en=utc(2026, 10, 23, 12, 0))
        resp = self.nuevo_pedido()
        self.assertEqual(resp.status_code, 400)
        self.assertIn('cerrado', str(resp.json()['punto_venta']))

    # --- DELETE / PATCH ---
    def test_no_se_borra_un_pedido_cobrado(self):
        resp = self.api.delete(f'/api/pedidos/{self.cobrado.id}/')
        self.assertEqual(resp.status_code, 409)
        self.assertTrue(Pedido.objects.filter(id=self.cobrado.id).exists())

    def test_si_se_borra_un_pedido_abierto(self):
        self.assertEqual(self.api.delete(f'/api/pedidos/{self.abierto.id}/').status_code, 204)
        self.assertFalse(Pedido.objects.filter(id=self.abierto.id).exists())

    def test_con_el_evento_cerrado_no_se_borra_ni_edita_nada(self):
        self.cerrar_expo()
        for pedido in (self.cobrado, self.abierto):
            self.assertEqual(self.api.delete(f'/api/pedidos/{pedido.id}/').status_code, 409)
            resp = self.api.patch(f'/api/pedidos/{pedido.id}/', {'lineas': [{'producto': self.agua.id, 'cantidad': 1}]}, format='json')
            self.assertEqual(resp.status_code, 409)
        self.assertEqual(Pedido.objects.count(), 2)
        self.assertEqual(self.abierto.lineas.get().producto_id, self.choripan.id)

    def test_se_puede_editar_un_pedido_abierto_de_un_evento_abierto(self):
        resp = self.api.patch(f'/api/pedidos/{self.abierto.id}/', {'lineas': [{'producto': self.agua.id, 'cantidad': 2}]}, format='json')
        self.assertEqual(resp.status_code, 200)

    # --- tablero ---
    def test_eventos_informa_el_cierre(self):
        self.cerrar_expo()
        data = {e['id']: e for e in self.api.get('/api/eventos/').json()}
        self.assertIsNotNone(data[self.expo.id]['cerrado_en'])
        self.assertIsNone(data[self.junio.id]['cerrado_en'])
        self.assertEqual(self.api.get('/api/resumen-ventas/', {'evento': self.expo.id}).json()['total_general'], 3000)


class PedidoAdminTests(DatosVentas):
    """Admin de pedidos: un pedido cobrado o de un evento cerrado no se borra, y los
    pedidos de un evento cerrado son de solo lectura (con sus inlines)."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        self.cobrado = self.pedido(self.expo, utc(2026, 10, 23, 0, 30))
        self.abierto = self.pedido(self.expo, utc(2026, 10, 23, 0, 40), cerrado=False)
        self.abierto2 = self.pedido(self.expo, utc(2026, 10, 23, 0, 50), cerrado=False)
        self.c = self.client_class()
        self.c.force_login(get_user_model().objects.create_superuser('admin', 'a@a.com', 'x'))

    def cerrar_expo(self):
        Evento.objects.filter(id=self.expo.id).update(activo=False, cerrado_en=utc(2026, 10, 24, 0, 0))

    def borrar_seleccionados(self, ids):
        return self.c.post('/admin/pdv/pedido/', {
            'action': 'delete_selected', '_selected_action': ids, 'post': 'yes'})

    def existen(self, *pedidos):
        return [Pedido.objects.filter(id=p.id).exists() for p in pedidos]

    # --- borrar ---
    def test_no_se_borra_un_pedido_cobrado(self):
        url = f'/admin/pdv/pedido/{self.cobrado.id}/delete/'
        self.assertEqual(self.c.get(url).status_code, 403)
        self.assertEqual(self.c.post(url, {'post': 'yes'}).status_code, 403)
        self.assertEqual(self.existen(self.cobrado), [True])

    def test_si_se_borra_un_pedido_abierto_de_un_evento_abierto(self):
        self.c.post(f'/admin/pdv/pedido/{self.abierto.id}/delete/', {'post': 'yes'})
        self.assertEqual(self.existen(self.abierto), [False])

    def test_eliminar_seleccionados_solo_abiertos_funciona(self):
        self.borrar_seleccionados([self.abierto.id, self.abierto2.id])
        self.assertEqual(self.existen(self.abierto, self.abierto2), [False, False])

    def test_eliminar_seleccionados_con_un_cobrado_no_borra_nada(self):
        self.borrar_seleccionados([self.abierto.id, self.cobrado.id])
        self.assertEqual(self.existen(self.abierto, self.cobrado), [True, True])

    def test_no_se_borra_nada_de_un_evento_cerrado(self):
        self.cerrar_expo()
        url = f'/admin/pdv/pedido/{self.abierto.id}/delete/'
        self.assertEqual(self.c.post(url, {'post': 'yes'}).status_code, 403)
        self.borrar_seleccionados([self.abierto.id, self.abierto2.id])
        self.assertEqual(self.existen(self.abierto, self.abierto2, self.cobrado), [True, True, True])

    def test_borrar_la_caja_no_arrastra_pedidos_protegidos(self):
        self.c.post(f'/admin/pdv/puntoventa/{self.caja.id}/delete/', {'post': 'yes'})
        self.assertEqual(self.existen(self.cobrado, self.abierto), [True, True])
        self.assertTrue(PuntoVenta.objects.filter(id=self.caja.id).exists())

    # --- solo lectura ---
    def test_pedido_de_evento_cerrado_es_de_solo_lectura(self):
        self.cerrar_expo()
        url = f'/admin/pdv/pedido/{self.abierto.id}/change/'
        resp = self.c.get(url)
        self.assertEqual(resp.status_code, 200)               # se puede ver
        self.assertContains(resp, 'Ver Pedido')
        self.assertNotContains(resp, 'name="lineas-0-cantidad"')  # sin campos editables en el inline
        self.assertNotContains(resp, 'name="total_final"')
        self.assertNotContains(resp, 'name="_save"')              # ni botón de guardar
        self.assertEqual(self.c.post(url, {'punto_venta': self.caja.id, 'evento': self.expo.id,
                                           'subtotal': '1', 'descuento_porcentaje': '0', 'total_impuestos': '0',
                                           'total_final': '1', 'veces_impreso': '0'}).status_code, 403)
        self.abierto.refresh_from_db()
        self.assertEqual(self.abierto.total_final, 1000)

    def test_pedido_de_evento_abierto_sigue_editable_con_sus_inlines(self):
        resp = self.c.get(f'/admin/pdv/pedido/{self.abierto.id}/change/')
        self.assertContains(resp, 'name="lineas-0-cantidad"')
        self.assertContains(resp, 'name="total_final"')
        self.assertContains(resp, 'name="_save"')
        self.assertContains(resp, 'pagos-TOTAL_FORMS')

    def test_no_se_puede_crear_un_pedido_en_un_evento_cerrado(self):
        self.cerrar_expo()
        resp = self.c.get('/admin/pdv/pedido/add/')
        self.assertNotContains(resp, 'Expo 2026')
        self.assertContains(resp, 'Peña 2026')


class CajerosTests(DatosVentas):
    """Roles y usuarios de caja (plan_permisos, ítem 1)."""

    CLAVE = 'Clave-Segura-2026!'

    def crear_cajeros(self):
        with mock.patch('pdv.management.commands.crear_cajeros.getpass', return_value=self.CLAVE):
            call_command('crear_cajeros', stdout=StringIO())

    def test_crear_cajeros_dos_veces_no_duplica(self):
        self.crear_cajeros()
        self.crear_cajeros()
        usuarios = User.objects.filter(username__in=['caja1', 'caja2', 'caja3'])
        self.assertEqual(usuarios.count(), 3)
        self.assertEqual(User.objects.count(), 3)
        self.assertEqual(Group.objects.filter(name='Cajero').count(), 1)
        for u in usuarios:
            self.assertEqual(list(u.groups.values_list('name', flat=True)), ['Cajero'])

    def test_cajero_no_es_staff_ni_entra_al_admin(self):
        self.crear_cajeros()
        caja1 = User.objects.get(username='caja1')
        self.assertFalse(caja1.is_staff)
        self.assertFalse(caja1.is_superuser)
        self.assertTrue(self.client.login(username='caja1', password=self.CLAVE))
        resp = self.client.get('/admin/')
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/admin/login/', resp['Location'])

    def test_dos_cajas_no_comparten_usuario(self):
        usuario = User.objects.create_user('caja1', password=self.CLAVE)
        self.caja.usuario = usuario
        self.caja.save()
        self.caja2.usuario = usuario
        with self.assertRaises(IntegrityError), transaction.atomic():
            self.caja2.save()


class PermisosApiTests(DatosVentas):
    """La API exige login y permisos por rol (plan_permisos, ítem 2)."""

    CLAVE = 'Clave-Segura-2026!'

    def setUp(self):
        self.admin = User.objects.create_superuser('admin', 'a@a.com', self.CLAVE)
        self.cajero = User.objects.create_user('caja1', password=self.CLAVE)
        self.cajero.groups.add(Group.objects.get_or_create(name='Cajero')[0])
        PuntoVenta.objects.filter(id=self.caja.id).update(usuario=self.cajero)
        self.sin_rol = User.objects.create_user('intruso', password=self.CLAVE)
        self.abierto = self.pedido(self.expo, utc(2026, 10, 23, 0, 40), cerrado=False)

    def cliente(self, usuario=None):
        c = APIClient()
        if usuario:
            c.force_authenticate(usuario)
        return c

    def endpoints_api(self):
        """(url, método) de cada ruta /api/ de urls.py, sin lista escrita a mano."""
        from django.urls import get_resolver
        resultado = []
        for patron in get_resolver().url_patterns:
            if not hasattr(patron, 'url_patterns'):
                continue  # admin u otras rutas que no son include('pdv.urls')
            for p in patron.url_patterns:
                ruta = str(p.pattern)
                if not ruta.startswith('api/'):
                    continue
                url = '/' + re.sub(r'<int:\w+>', str(self.abierto.id), ruta)
                for metodo in p.callback.cls().allowed_methods:
                    if metodo not in ('OPTIONS', 'HEAD'):
                        resultado.append((url, metodo))
        return resultado

    def llamar(self, cliente, url, metodo):
        return getattr(cliente, metodo.lower())(url, {}, format='json')

    def pedido_nuevo(self):
        return {'punto_venta': self.caja.id, 'lineas': [{'producto': self.choripan.id, 'cantidad': 1}]}

    def test_2_1_sin_login_todos_los_endpoints_rechazan(self):
        endpoints = self.endpoints_api()
        self.assertGreaterEqual(len(endpoints), 18)
        for url, metodo in endpoints:
            resp = self.llamar(self.cliente(), url, metodo)
            self.assertIn(resp.status_code, (401, 403), f'{metodo} {url} -> {resp.status_code}')

    def test_2_2_cajero_crea_cobra_y_reimprime(self):
        c = self.cliente(self.cajero)
        resp = c.post('/api/pedidos/', self.pedido_nuevo(), format='json')
        self.assertEqual(resp.status_code, 201)
        cobro = dict(self.pedido_nuevo(), pagos=[{'metodo': 'EF', 'monto': '1000'}])
        self.assertEqual(c.post('/api/pedidos/', cobro, format='json').status_code, 201)
        with mock.patch('pdv.utils.ticket_handler.imprimir_ticket'):
            resp = c.post(f'/api/pedidos/{self.abierto.id}/reimprimir/', {}, format='json')
        self.assertEqual(resp.status_code, 200)
        for url in ('/api/pedidos/abiertos/', '/api/pedidos/historial/', '/api/productos/', '/api/pdv/',
                    f'/api/pedidos/{self.abierto.id}/'):
            self.assertEqual(c.get(url).status_code, 200, url)

    def test_2_3_cajero_recibe_403_en_lo_del_administrador(self):
        c = self.cliente(self.cajero)
        for url in ('/api/resumen-ventas/', '/api/resumen-ventas/xlsx/', '/api/resumen-ventas/imprimir/',
                    '/api/eventos/', f'/api/productos/{self.choripan.id}/'):
            self.assertEqual(c.get(url).status_code, 403, url)
        self.assertEqual(c.post('/api/cierre-caja/', {}, format='json').status_code, 403)
        self.assertEqual(c.post('/api/test-print/', {}, format='json').status_code, 403)
        resp = c.post('/api/productos/', {'nombre': 'X', 'precio': '1', 'categoria': 1}, format='json')
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(c.patch(f'/api/productos/{self.choripan.id}/', {'precio': '1'}, format='json').status_code, 403)

    def test_usuario_sin_rol_no_entra_a_nada(self):
        c = self.cliente(self.sin_rol)
        for url, metodo in self.endpoints_api():
            self.assertEqual(self.llamar(c, url, metodo).status_code, 403, f'{metodo} {url}')

    def test_2_4_administrador_puede_usar_todo(self):
        c = self.cliente(self.admin)
        for url in ('/api/resumen-ventas/', '/api/resumen-ventas/xlsx/', '/api/resumen-ventas/imprimir/',
                    '/api/eventos/', '/api/productos/', f'/api/productos/{self.choripan.id}/',
                    '/api/pedidos/abiertos/', '/api/pedidos/historial/', '/api/pdv/'):
            self.assertEqual(c.get(url).status_code, 200, url)
        self.assertEqual(c.post('/api/cierre-caja/', {}, format='json').status_code, 200)
        self.assertEqual(c.post('/api/pedidos/', self.pedido_nuevo(), format='json').status_code, 201)

    def test_2_5_post_sin_csrf_con_sesion_se_rechaza(self):
        c = APIClient(enforce_csrf_checks=True)
        self.assertTrue(c.login(username='caja1', password=self.CLAVE))
        antes = Pedido.objects.count()
        resp = c.post('/api/pedidos/', self.pedido_nuevo(), format='json')
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(Pedido.objects.count(), antes)
        # con el token, la misma llamada entra
        c.get('/')  # ensure_csrf_cookie
        token = c.cookies['csrftoken'].value
        resp = c.post('/api/pedidos/', self.pedido_nuevo(), format='json', HTTP_X_CSRFTOKEN=token)
        self.assertEqual(resp.status_code, 201)


class CajaPropiaTests(DatosVentas):
    """Cada cajero opera solo sobre su caja (plan_permisos, ítem 3)."""

    CLAVE = 'Clave-Segura-2026!'

    def setUp(self):
        grupo = Group.objects.get_or_create(name='Cajero')[0]
        self.u1 = User.objects.create_user('caja1', password=self.CLAVE)
        self.u2 = User.objects.create_user('caja2', password=self.CLAVE)
        self.u3 = User.objects.create_user('caja3', password=self.CLAVE)   # cajero sin caja asignada
        for u in (self.u1, self.u2, self.u3):
            u.groups.add(grupo)
        PuntoVenta.objects.filter(id=self.caja.id).update(usuario=self.u1)
        PuntoVenta.objects.filter(id=self.caja2.id).update(usuario=self.u2)
        self.admin = User.objects.create_superuser('admin', 'a@a.com', self.CLAVE)
        self.abierto1 = self.pedido(self.expo, utc(2026, 10, 23, 0, 10), cerrado=False, caja=self.caja)
        self.abierto2 = self.pedido(self.expo, utc(2026, 10, 23, 0, 20), cerrado=False, caja=self.caja2)
        self.cobrado1 = self.pedido(self.expo, utc(2026, 10, 23, 0, 30), caja=self.caja)
        self.cobrado2 = self.pedido(self.expo, utc(2026, 10, 23, 0, 40), caja=self.caja2)

    def api(self, usuario):
        c = APIClient()
        c.force_authenticate(usuario)
        return c

    def nuevo(self, punto_venta):
        return {'punto_venta': punto_venta.id, 'lineas': [{'producto': self.choripan.id, 'cantidad': 1}]}

    def ids(self, resp):
        datos = resp.json()
        return {p['id'] for p in (datos['resultados'] if isinstance(datos, dict) else datos)}

    def test_3_1_la_caja_sale_del_usuario_no_de_lo_que_manda_el_cliente(self):
        resp = self.api(self.u2).post('/api/pedidos/', self.nuevo(self.caja), format='json')   # caja2 dice ser caja1
        self.assertEqual(resp.status_code, 201)
        self.assertEqual(Pedido.objects.get(id=resp.json()['id']).punto_venta_id, self.caja2.id)

    def test_3_1_sin_punto_venta_en_el_cuerpo_tambien_usa_la_del_usuario(self):
        cuerpo = self.nuevo(self.caja)
        del cuerpo['punto_venta']
        resp = self.api(self.u1).post('/api/pedidos/', cuerpo, format='json')
        self.assertEqual(resp.status_code, 201)
        self.assertEqual(Pedido.objects.get(id=resp.json()['id']).punto_venta_id, self.caja.id)

    def test_cuerpo_que_no_es_un_objeto_da_400_y_no_500(self):
        resp = self.api(self.u1).post('/api/pedidos/', [1, 2], format='json')
        self.assertEqual(resp.status_code, 400)

    def test_el_administrador_elige_la_caja(self):
        resp = self.api(self.admin).post('/api/pedidos/', self.nuevo(self.caja2), format='json')
        self.assertEqual(resp.status_code, 201)
        self.assertEqual(Pedido.objects.get(id=resp.json()['id']).punto_venta_id, self.caja2.id)

    def test_3_2_el_cajero_no_ve_el_historial_ni_los_abiertos_de_otra_caja(self):
        c = self.api(self.u1)
        for url in ('/api/pedidos/historial/', '/api/pedidos/abiertos/'):
            propios = self.ids(c.get(url))
            self.assertTrue(propios, url)
            self.assertTrue(propios <= {self.abierto1.id, self.cobrado1.id}, url)
            # pedir explícitamente la otra caja no cambia nada
            self.assertEqual(self.ids(c.get(url, {'punto_venta_id': self.caja2.id})), propios, url)
        self.assertEqual(self.ids(self.api(self.admin).get('/api/pedidos/historial/')),
                         {self.abierto1.id, self.abierto2.id, self.cobrado1.id, self.cobrado2.id})

    def test_el_cajero_ve_solo_su_caja_en_la_lista_de_cajas(self):
        resp = self.api(self.u1).get('/api/pdv/')
        self.assertEqual([p['id'] for p in resp.json()], [self.caja.id])
        self.assertEqual(len(self.api(self.admin).get('/api/pdv/').json()), 2)

    def test_3_3_el_cajero_no_toca_pedidos_de_otra_caja(self):
        c = self.api(self.u1)
        otro = self.abierto2.id
        self.assertEqual(c.delete(f'/api/pedidos/{otro}/').status_code, 403)
        self.assertTrue(Pedido.objects.filter(id=otro).exists())
        self.assertEqual(c.get(f'/api/pedidos/{otro}/').status_code, 403)
        self.assertEqual(c.patch(f'/api/pedidos/{otro}/', {'lineas': [{'producto': self.agua.id, 'cantidad': 1}]},
                                 format='json').status_code, 403)
        self.assertEqual(c.post(f'/api/pedidos/{otro}/reimprimir/', {}, format='json').status_code, 403)
        self.assertEqual(c.post(f'/api/pedidos/{otro}/imprimir-local/', {}, format='json').status_code, 403)
        self.assertEqual(c.get(f'/api/pedidos/{otro}/imprimir-pdf/').status_code, 403)
        self.assertEqual(c.get(f'/api/pedidos/{otro}/ticket-html/').status_code, 403)

    def test_el_cajero_si_borra_un_pedido_abierto_de_su_caja(self):
        self.assertEqual(self.api(self.u1).delete(f'/api/pedidos/{self.abierto1.id}/').status_code, 204)

    def test_el_administrador_toca_pedidos_de_cualquier_caja(self):
        self.assertEqual(self.api(self.admin).delete(f'/api/pedidos/{self.abierto2.id}/').status_code, 204)

    def test_3_4_cajero_sin_caja_no_puede_operar(self):
        c = self.api(self.u3)
        resp = c.post('/api/pedidos/', self.nuevo(self.caja), format='json')
        self.assertEqual(resp.status_code, 403)
        self.assertIn('caja', resp.json()['detail'].lower())
        self.assertFalse(Pedido.objects.filter(creado__gte=utc(2026, 10, 24, 0, 0)).exists())
        for url in ('/api/pedidos/historial/', '/api/pedidos/abiertos/', '/api/pdv/', f'/api/pedidos/{self.abierto1.id}/'):
            self.assertEqual(c.get(url).status_code, 403, url)

    def test_3_5_siguen_las_reglas_de_pedidos_cobrados_y_eventos_cerrados(self):
        c = self.api(self.u1)
        self.assertEqual(c.delete(f'/api/pedidos/{self.cobrado1.id}/').status_code, 409)
        self.assertTrue(Pedido.objects.filter(id=self.cobrado1.id).exists())
        Evento.objects.filter(id=self.expo.id).update(activo=False, cerrado_en=utc(2026, 10, 23, 12, 0))
        self.assertEqual(c.delete(f'/api/pedidos/{self.abierto1.id}/').status_code, 409)
        resp = c.patch(f'/api/pedidos/{self.abierto1.id}/', {'lineas': [{'producto': self.agua.id, 'cantidad': 1}]}, format='json')
        self.assertEqual(resp.status_code, 409)
        self.assertEqual(c.post('/api/pedidos/', self.nuevo(self.caja), format='json').status_code, 400)
